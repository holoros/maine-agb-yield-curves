"""Build the Maine FIA AGB yield curve workbook from Cardinal aggregate outputs (no plot identifiers, no coordinates).
Usage: python3 build_xlsx.py <results_dir with bakuzis/ curves/ stress/> <out.xlsx>"""
import sys, csv
from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Font, Border, Side
from openpyxl.utils import get_column_letter

R, out = sys.argv[1], sys.argv[2]
def rd(f):
    with open(f"{R}/{f}") as fh: return list(csv.DictReader(fh))
FT = {"SF": "Spruce-fir", "MBB": "Maple-beech-birch", "AB": "Aspen-birch", "WRPH": "White-red pine-hemlock",
      "OAK": "Oak and oak-pine", "EAC": "Elm-ash-cottonwood"}
SUP = InlineFont(vertAlign="superscript", b=True); B = InlineFont(b=True)
def hdr_cell(txt):
    if "^" not in txt: return txt
    parts, s = [], txt
    while "^" in s:
        pre, rest = s.split("^", 1); exp = ""
        while rest and rest[0] in "-0123456789": exp += rest[0]; rest = rest[1:]
        parts += [TextBlock(B, pre), TextBlock(SUP, exp.replace("-", "−"))]; s = rest
    if s: parts.append(TextBlock(B, s))
    return CellRichText(parts)
thin = Side(style="thin"); ctr = Alignment(horizontal="center", vertical="center", wrap_text=True)
def num(x, d=3):
    try:
        v = float(x)
        return round(v, d)
    except Exception:
        return x
def sheet(wb, name, hdr, rows, widths=None):
    ws = wb.create_sheet(name)
    for j, h in enumerate(hdr, 1):
        c = ws.cell(1, j, hdr_cell(h)); c.font = Font(bold=True); c.alignment = ctr; c.border = Border(top=thin, bottom=thin)
    for i, r in enumerate(rows, 2):
        for j, v in enumerate(r, 1): ws.cell(i, j, v).alignment = ctr
    last = len(rows) + 1
    for j in range(1, len(hdr) + 1):
        if last > 1: ws.cell(last, j).border = Border(bottom=thin)
        ws.column_dimensions[get_column_letter(j)].width = widths[j - 1] if widths else max(11, min(26, len(str(hdr[j - 1])) + 2))
    ws.row_dimensions[1].height = 42; ws.freeze_panes = "A2"
    return ws

wb = Workbook(); wb.remove(wb.active)
pick = rd("stress/pick_DATA.csv"); par = rd("stress/final_params_DATA.csv"); lev = rd("stress/final_level_DATA.csv")
fc = rd("stress/final_curves_capped_DATA.csv"); sc = rd("stress/scoring_DATA.csv"); bs = rd("stress/bootstrap_decline_DATA.csv")
lo = rd("stress/loso_DATA.csv"); mg = rd("bakuzis/magnitude_reference_DATA.csv")
notes = [
 "Maine FIA generic aboveground live tree biomass (AGB) yield curves by forest type group and ecoregion section (CRSF, October 15, 2026; DRAFT).",
 "Data: FIADB Maine annual inventory 1999 to 2024, single condition forested plots; harvest free plot visits and every harvest free remeasurement interval (consecutive and longer, up to 25 yr).",
 "Curve: AGB(t) = min(cap, exp(mu_type + u_cell) x f(t)), with f the chosen decline-capable form for the type (Forms sheet) and cap the type's observed 99th percentile AGB.",
 "Forms: CRS f = (1-exp(-k t))^p exp(-d t); CRD f = (1-exp(-k t))^p (1 - q (1-exp(-d t))); HG f = t^b exp(-c t); CR (reference) f = (1-exp(-k t))^p.",
 "Strategies: S1 remeasurement only (algebraic difference on intervals); S2 joint chronosequence plus repeated measures mixed model (profile likelihood).",
 "AGB median: log scale back transform (typical plot). AGB mean: median x exp((s_plot^2 + s_res^2)/2). Lower and upper 95%: plot cluster bootstrap confidence interval on the median curve. PI90: 90% prediction interval for a single plot.",
 "Extrapolated = TRUE beyond the cell's 99th percentile observed stand age. Capped = TRUE where the uncapped median exceeds the cap. Treat both as assumptions, not estimates.",
 "Ecoregion sections from the FIA ECOSUBCD section code. No plot identifiers or coordinates are included in this workbook.",
 "Decline: the decline-capable forms let the data decide; the Decline sheet gives the bootstrap probability that a curve peaks before 250 yr."]
ws = wb.create_sheet("README")
for i, t in enumerate(notes, 1): ws.cell(i, 1, t).alignment = Alignment(wrap_text=True, vertical="top")
ws.column_dimensions["A"].width = 160
sheet(wb, "Forms", ["Type code", "Forest type group", "Strategy", "Form", "Basis"],
      [[p["ftg"], FT[p["ftg"]], p["strat"], p["form"], p["basis"]] for p in pick], [10, 26, 10, 8, 22])
sheet(wb, "Parameters", ["Type code", "Strategy", "Form", "Parameter", "Estimate", "mu (log scale level)"],
      [[p["ftg"], p["strat"], p["form"], p["param"], num(p["est"], 6), num(p["mu"], 4)] for p in par])
sheet(wb, "Level", ["Cell (type:section)", "u_cell (log scale)", "s_plot", "s_cell", "s_res"],
      [[l["cell"], num(l["u_cell"], 4), num(l["s_plot"], 4), num(l["s_cell"], 4), num(l["s_res"], 4)] for l in lev], [20, 14, 10, 10, 10])
caps = {m["ftg"]: m for m in mg}
sheet(wb, "Caps", ["Type code", "Observed median AGB (Mg ha^-1)", "Cap: observed 99th percentile AGB (Mg ha^-1)", "Observed 99.9th percentile AGB (Mg ha^-1)", "Observed 99th percentile stand age (yr)"],
      [[k, num(caps[k]["AGB_q50"], 1), num(caps[k]["AGB_q99"], 1), num(caps[k]["AGB_q999"], 1), num(caps[k]["age_q99"], 0)] for k in FT if k in caps])
H = ["Type code", "Section", "Stand age (yr)", "AGB median (Mg ha^-1)", "AGB mean (Mg ha^-1)", "Lower 95% (Mg ha^-1)", "Upper 95% (Mg ha^-1)",
     "PI90 lower (Mg ha^-1)", "PI90 upper (Mg ha^-1)", "Capped", "Extrapolated", "Plots in cell"]
row = lambda r: [r["ftg"], r["section"], int(float(r["age"])), num(r["AGB_med"], 1), num(r["AGB_mean"], 1), num(r["lo95"], 1), num(r["hi95"], 1),
                 num(r["PI90_lo"], 1), num(r["PI90_hi"], 1), r["capped"], r["extrapolated"], r["n_plots"]]
sheet(wb, "Curves_5yr", H, [row(r) for r in fc if int(float(r["age"])) % 5 == 0])
sheet(wb, "Curves_annual", H, [row(r) for r in fc])
sheet(wb, "Decline", ["Type code", "Bootstrap draws", "P(peak before 250 yr)", "Peak age q05 (yr)", "Peak age q50 (yr)", "Peak age q95 (yr)",
      "B(250)/max q05", "B(250)/max q50", "B(250)/max q95"],
      [[b["ftg"], b["nboot"], num(b["p_decline_250"]), num(b["peak_q05"], 0), num(b["peak_q50"], 0), num(b["peak_q95"], 0),
        num(b["r250_q05"]), num(b["r250_q50"]), num(b["r250_q95"])] for b in bs])
sheet(wb, "Scoring", ["Type code", "Strategy", "Form", "Passes screens", "Peak age (yr)", "B(250)/max", "Generic CV RMSE (Mg ha^-1)", "LOSO RMSE (Mg ha^-1)",
      "LOSO PI90 coverage", "Stress spread B(200)/B(100)", "Score"],
      [[s["ftg"], s["strat"], s["form"], s["pass"], s["peak_age"], num(s["r250"]), num(s["cv_rmse"], 1), num(s["loso_rmse"], 1), num(s["loso_cover90"]),
        num(s["spread_g200_100"]), s.get("score", "")] for s in sc])
import os
if os.path.exists(f"{R}/benchmark/benchmark_by_type_DATA.csv"):
    bm = rd("benchmark/benchmark_by_type_DATA.csv")
    sheet(wb, "Benchmark_prior", ["Type code", "Visits", "Observed mean AGB (Mg ha^-1)", "Prior ME.AGB bias (Mg ha^-1)", "Prior RMSE (Mg ha^-1)",
          "New mean curve bias (Mg ha^-1)", "New RMSE, in sample (Mg ha^-1)", "New RMSE, cross validated (Mg ha^-1)"],
          [[b["ftg"], b["n"], num(b["obs_mean"], 1), num(b["prior_bias"], 1), num(b["prior_rmse"], 1), num(b["new_bias_mean"], 1), num(b["new_rmse_mean"], 1), num(b.get("new_cv_rmse", ""), 1)] for b in bm])
wb.save(out); print("wrote", out)
